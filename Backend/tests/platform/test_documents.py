"""Generate real files locally; exercise extraction and upload orchestration."""
from types import SimpleNamespace
from unittest.mock import MagicMock
from uuid import uuid4
import pytest
from app.utils.text_extractor import extract_text_from_file, _sanitize
from app.utils.text_chunker import chunk_text
from app.utils.question_parser import parse_testbank_text_to_questions


@pytest.fixture(params=["txt", "pdf", "docx", "pptx"])
def source_file(request, tmp_path):
    extension = request.param
    path = tmp_path / f"lesson.{extension}"
    phrase = "The product rule differentiates a product of two functions."
    if extension == "txt":
        path.write_text(phrase)
    elif extension == "pdf":
        import fitz
        with fitz.open() as doc:
            doc.new_page().insert_text((72, 72), phrase)
            doc.save(path)
    elif extension == "docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph(phrase)
        doc.save(path)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = phrase
        doc.save(path)
    return path, extension, phrase


def test_real_file_extraction(source_file):
    path, extension, phrase = source_file
    result = extract_text_from_file(str(path), extension)
    assert phrase in result["text"]
    assert result["metadata"]


def test_unsupported_file_rejected(tmp_path):
    with pytest.raises(ValueError, match="Unsupported"):
        extract_text_from_file(str(tmp_path / "data.exe"), "exe")


def test_sanitization_preserves_math_and_paragraphs():
    data = {"text": "\x00Formula $x^2$\nExplanation\tend", "pages": ["bad\x08text"]}
    assert _sanitize(data) == {"text": "Formula $x^2$\nExplanation\tend", "pages": ["badtext"]}


def test_parser_outage_falls_back_to_local_extraction(monkeypatch, tmp_path):
    from docx import Document
    from app.utils import text_extractor as extractor
    path = tmp_path / "testbank.docx"
    doc = Document()
    doc.add_paragraph("1. What is two plus two?")
    doc.save(path)
    monkeypatch.setattr(extractor, "extract_text_with_llamaparse", MagicMock(side_effect=RuntimeError("provider unavailable")))
    assert "two plus two" in extractor.extract_text_from_file(str(path), "docx", is_testbank=True)["text"]


def test_chunk_boundaries_cover_source_and_overlap():
    source = "abcdefghijklmnopqrstuvwxyz" * 100
    chunks = chunk_text(source, chunk_size=100, overlap=20)
    assert chunks[0]["start"] == 0
    assert chunks[-1]["end"] == len(source)
    for index, chunk in enumerate(chunks):
        assert chunk["index"] == index
        assert chunk["text"] == source[chunk["start"]:chunk["end"]]
        if index:
            assert chunks[index - 1]["end"] - chunk["start"] == 20


@pytest.mark.parametrize("source", ["", "   ", "\n\t"])
def test_empty_document_has_no_chunks(source):
    assert chunk_text(source) == []


def test_testbank_extracts_options_and_math_answer():
    questions = parse_testbank_text_to_questions(
        "1. Differentiate $x^2$.\nA. $x$\nB. $2x$\nAnswer: B\n\n2. What is 1+1?\nA. 2\nB. 3\nAnswer: A", uuid4())
    assert len(questions) == 2
    assert questions[0]["options"]["B"] == "$2x$"
    assert questions[0].get("correct_option_id", questions[0].get("correct_answer")) == "B"


def test_upload_runs_extraction_chunking_and_embedding(db_session, test_module, teacher_user, tmp_path, monkeypatch):
    from app.services import document as service
    from app.models.document_chunk import DocumentChunk
    path = tmp_path / "lesson.txt"
    path.write_text("The product rule. " * 100)
    monkeypatch.setattr(service.storage_service, "upload_file", MagicMock(return_value="https://test.invalid/lesson.txt"))
    monkeypatch.setattr(service.storage_service, "download_file_temporarily", MagicMock(return_value=str(path)))
    embed = MagicMock(return_value=2)
    monkeypatch.setattr(service, "generate_embeddings_for_document", embed)
    result = service.handle_document_upload(db_session, path.read_bytes(), "lesson.txt", teacher_user.id,
                                            module_name=test_module.name)
    assert result.processing_status == "embedded"
    assert db_session.query(DocumentChunk).filter_by(document_id=result.id).count() > 0
    embed.assert_called_once()
    assert not path.exists(), "temporary download must be cleaned up"


def test_storage_failure_does_not_create_document(db_session, test_module, teacher_user, monkeypatch):
    from app.services import document as service
    from app.models.document import Document
    from fastapi import HTTPException
    monkeypatch.setattr(service.storage_service, "upload_file", MagicMock(side_effect=RuntimeError("storage unavailable")))
    with pytest.raises(HTTPException) as error:
        service.handle_document_upload(db_session, b"Lesson", "lesson.txt", teacher_user.id, module_name=test_module.name)
    assert error.value.status_code == 500
    assert db_session.query(Document).count() == 0
